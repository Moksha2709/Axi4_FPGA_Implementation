"""Builds E:\\FPGA_Project\\docs\\AXI4_Viva_Presentation.pptx

Run:  python docs\\build_deck.py        (needs python-pptx)
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "AXI4_Viva_Presentation.pptx")

# ---- palette: same colour roles as the block diagram -----------------------
NAVY, LIGHT, WHITE = "172033", "F5F7FA", "FFFFFF"
INK, MUTED, RULE = "1E2A3B", "4F5F74", "D5DCE6"
AMBER, AMBER_D, AMBER_F = "FFC845", "8F5C00", "FFE9A8"
BLUE, BLUE_F = "3B6FD0", "D6E6FF"
VIOLET, VIOLET_F = "7152C9", "E6DEFA"
GREEN, GREEN_F = "2E8B57", "D9F0E3"
ON_DARK, ON_DARK_M = "EEF2F8", "B7C3D6"

HEAD, BODY, MONO = "Cambria", "Calibri", "Consolas"
W, H = 13.333, 7.5
MX = 0.7


def rgb(h):
    return RGBColor.from_string(h)


prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
BLANK = prs.slide_layouts[6]
page_no = [0]


# ---- helpers ---------------------------------------------------------------
def new_slide(bg=LIGHT, notes=""):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(bg)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    page_no[0] += 1
    return s


def text(slide, x, y, w, h, paras, size=16, bold=False, color=INK, font=BODY,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False, name=None):
    """paras: str | list[str | dict(text,size,bold,color,font,after,italic,align)]"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [paras]
    for i, p in enumerate(paras):
        d = {"text": p} if isinstance(p, str) else p
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = d.get("align", align)
        run = para.add_run()
        run.text = d["text"]
        f = run.font
        f.size = Pt(d.get("size", size))
        f.bold = d.get("bold", bold)
        f.italic = d.get("italic", italic)
        f.name = d.get("font", font)
        f.color.rgb = rgb(d.get("color", color))
        if "after" in d:
            para.space_after = Pt(d["after"])
    return tb


def box(slide, x, y, w, h, fill=WHITE, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        radius=0.06, name=None):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        sp.name = name
    sp.fill.solid()
    sp.fill.fore_color.rgb = rgb(fill)
    if line:
        sp.line.color.rgb = rgb(line)
        sp.line.width = Pt(1.5)
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sp.adjustments[0] = radius
    sp.text_frame.text = ""
    return sp


def pill(slide, x, y, w, h, label, fill=AMBER_F, color=INK, size=14, bold=True, font=BODY):
    sp = box(slide, x, y, w, h, fill=fill, radius=0.5)
    tf = sp.text_frame
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = font
    r.font.color.rgb = rgb(color)
    return sp


def arrow(slide, x1, y1, x2, y2, color=INK, width=2.25, head=True):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                   Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    if head:
        ln = c.line._get_or_add_ln()
        tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
        ln.append(tail)
    return c


def line(slide, x1, y1, x2, y2, color=INK, width=2.25):
    return arrow(slide, x1, y1, x2, y2, color, width, head=False)


def title(slide, label, dark=False):
    text(slide, MX, 0.45, W - 2 * MX, 0.9, label, size=38, bold=True,
         color=ON_DARK if dark else INK, font=HEAD, anchor=MSO_ANCHOR.MIDDLE, name="Title")


def footer(slide, dark=False):
    c = ON_DARK_M if dark else MUTED
    text(slide, MX, 7.0, 10, 0.3, "AXI4: input D flip-flops, a master, a slave and output D flip-flops, in Verilog",
         size=11, color=c, name="Footer")
    text(slide, W - MX - 1.0, 7.0, 1.0, 0.3, str(page_no[0]), size=11, color=c,
         align=PP_ALIGN.RIGHT, name="Page number")


def picture(slide, path, x, y, w=None, h=None, alt=""):
    p = slide.shapes.add_picture(path, Inches(x), Inches(y),
                                 Inches(w) if w else None, Inches(h) if h else None)
    p._element.nvPicPr.cNvPr.set("descr", alt)
    return p


def card_row(slide, y, h, cards, n, head_size=22, body_size=16, gap=0.22):
    """n equal cards across the content width. cards: (fill, heading, sub, body)"""
    cw = (W - 2 * MX - (n - 1) * gap) / n
    x = MX
    for fill, head, sub, body in cards:
        box(slide, x, y, cw, h, fill=fill)
        text(slide, x + 0.28, y + 0.25, cw - 0.56, 0.5, head, size=head_size, bold=True, font=HEAD)
        yy = y + 0.85
        if sub:
            text(slide, x + 0.28, yy, cw - 0.56, 0.35, sub, size=13, color=MUTED, font=MONO)
            yy += 0.5
        text(slide, x + 0.28, yy, cw - 0.56, h - (yy - y) - 0.2, body, size=body_size)
        x += cw + gap


# =========================================================================
# 1  Cover
# =========================================================================
s = new_slide(NAVY, "Open with the one-sentence summary: we built an AXI4 master and an AXI4 memory in "
              "Verilog. The testbench loads an address and four words into input D flip-flops. The "
              "master takes them from there, writes them to the slave over the five AXI channels, then "
              "reads them back into output D flip-flops. The master and the slave are connected "
              "directly. Everything shown comes from the simulation in the project folder, so every "
              "number can be reproduced live. The five amber pills are the five AXI channels; we come "
              "back to each one.")
text(s, 0.9, 1.5, 11, 0.4, "VIVA DEMONSTRATION", size=18, bold=True, color=AMBER, name="Eyebrow")
text(s, 0.9, 2.05, 11.4, 2.4, "AXI4: five channels, a master and a memory", size=60, bold=True,
     color=ON_DARK, font=HEAD, name="Title")
text(s, 0.9, 4.25, 10.8, 1.0,
     "Input D flip-flops, master, slave, output D flip-flops: a write and a read, "
     "and 76 checks pass in simulation.",
     size=24, color=ON_DARK_M)
px = 0.9
for lab, w_ in (("AW", 0.9), ("W", 0.8), ("B", 0.8), ("AR", 0.9), ("R", 0.8)):
    pill(s, px, 5.55, w_, 0.5, lab, fill=AMBER, color=NAVY, size=18)
    px += w_ + 0.18
text(s, 0.9, 6.7, 11, 0.4, "[Your name]   ·   [Course and institution]   ·   [Date]",
     size=16, color=ON_DARK_M, name="Presenter")

# =========================================================================
# 2  What we built
# =========================================================================
s = new_slide(LIGHT, "Five small files, and the top level that wires them. The input D flip-flops hold "
              "the address and four words the testbench loads. The master is the active side: it takes "
              "them from there and starts the AXI transfers when start goes high. The slave is a 4 KiB "
              "memory that answers the five channels. The output D flip-flops catch the words that "
              "come back. The top level only wires them together; the master and slave are connected "
              "directly.")
title(s, "What we built")
text(s, MX, 1.4, 11.9, 0.5, "Five files, each with one job.", size=20, color=MUTED)
card_row(s, 2.2, 3.9, [
    (AMBER_F, "Input DFFs", "axi4_in_dff.v",
     "Five registers: an address and four words. The testbench loads them."),
    (BLUE_F, "Master", "axi4_master.v",
     "Takes the data from the input DFFs. Writes it, then reads it back."),
    (VIOLET_F, "Slave", "axi4_slave.v",
     "A 4 KiB memory. Answers bursts of up to 256 beats."),
    (AMBER_F, "Output DFFs", "axi4_out_dff.v",
     "Four registers that hold the words read back."),
    (GREEN_F, "Top", "axi4_top.v",
     "Wires it all together. Master and slave connect directly."),
], 5, head_size=20, body_size=15, gap=0.18)
text(s, MX, 6.4, 11.9, 0.4,
     "Full AXI4 rather than AXI4-Lite: bursts, transaction IDs, WLAST and RLAST.",
     size=16, color=MUTED, italic=True)
footer(s)

# =========================================================================
# 3  Five channels
# =========================================================================
s = new_slide(LIGHT, "AXI separates a transaction into five independent one-way channels, so a read "
              "can proceed while a write waits. A write uses AW for the address, W for the data beats, "
              "and B for one response at the end. A read uses AR for the address and R for the data "
              "coming back. Notice B and R point the other way.")
title(s, "Five independent channels")
box(s, 0.9, 1.65, 2.5, 4.95, fill=BLUE_F, line=BLUE)
text(s, 0.9, 3.4, 2.5, 0.8, [{"text": "Master", "size": 24, "bold": True, "font": HEAD,
                                "align": PP_ALIGN.CENTER},
                               {"text": "axi4_master", "size": 14, "font": MONO, "color": MUTED,
                                "align": PP_ALIGN.CENTER}])
box(s, 9.93, 1.65, 2.5, 4.95, fill=VIOLET_F, line=VIOLET)
text(s, 9.93, 3.4, 2.5, 0.8, [{"text": "Slave", "size": 24, "bold": True, "font": HEAD,
                                 "align": PP_ALIGN.CENTER},
                                {"text": "axi4_slave", "size": 14, "font": MONO, "color": MUTED,
                                 "align": PP_ALIGN.CENTER}])
text(s, 3.7, 1.7, 2, 0.3, "WRITE", size=12, bold=True, color=MUTED)
text(s, 3.7, 4.55, 2, 0.3, "READ", size=12, bold=True, color=MUTED)
rows = [
    (2.4, "AW   write address", True), (3.25, "W   write data", True),
    (4.1, "B   write response", False),
    (5.2, "AR   read address", True), (6.05, "R   read data", False),
]
for y, lab, fwd in rows:
    if fwd:
        arrow(s, 3.45, y, 9.88, y, color=INK)
    else:
        arrow(s, 9.88, y, 3.45, y, color=INK)
    pill(s, 5.25, y - 0.24, 3.3, 0.48, lab, fill=AMBER_F, size=16)
footer(s)

# =========================================================================
# 4  Handshake
# =========================================================================
s = new_slide(LIGHT, "Every channel uses the same handshake. The sender raises VALID when its data is "
              "real. The receiver raises READY when it can take it. The transfer happens in the one "
              "clock cycle where both are high. Two rules keep it safe. First, once VALID is up the "
              "data must not change until the handshake completes. Second, a sender may not wait for "
              "READY before raising VALID, or two sides could wait on each other forever. Our "
              "testbench checks both rules on every run.")
title(s, "The VALID and READY handshake")
box(s, MX, 1.65, 5.6, 2.2, fill=WHITE)
text(s, MX + 0.3, 1.85, 5.0, 0.4, "Transfer", size=20, bold=True, font=HEAD)
text(s, MX + 0.3, 2.35, 5.0, 1.4,
     "The sender raises VALID when its data is real. The receiver raises READY when it can take it. "
     "The transfer happens in the cycle where both are high.", size=16)
box(s, MX, 4.05, 5.6, 2.6, fill=WHITE)
text(s, MX + 0.3, 4.25, 5.0, 0.4, "Two rules", size=20, bold=True, font=HEAD)
text(s, MX + 0.3, 4.75, 5.0, 1.8, [
    {"text": "Once VALID is up, the data and VALID stay put until the handshake completes.", "after": 8},
    {"text": "A sender may not wait for READY before raising VALID.", "after": 0}], size=16)
wx, wy, cw = 6.9, 1.9, 0.7
box(s, 6.5, 1.65, 6.13, 5.0, fill=WHITE)
text(s, 6.8, 1.8, 3, 0.35, "One transfer, cycle by cycle", size=14, color=MUTED)
sigs = [("CLK", None), ("VALID", [0, 1, 1, 1, 0, 0]), ("READY", [0, 0, 0, 1, 1, 0])]
base = wy + 1.35
for i in range(6):
    xx = wx + 0.85 + i * cw
    line(s, xx, wy + 0.4, xx, wy + 3.75, color=RULE, width=1)
    text(s, xx, wy + 0.1, cw, 0.3, str(i), size=12, color=MUTED, align=PP_ALIGN.CENTER)
for r_, (nm, bits) in enumerate(sigs):
    yb = base + r_ * 0.9
    yt = yb - 0.5
    text(s, 6.8, yb - 0.45, 0.9, 0.4, nm, size=14, bold=True, font=MONO)
    if bits is None:
        for i in range(6):
            x0 = wx + 0.85 + i * cw
            line(s, x0, yt, x0 + cw / 2, yt, BLUE)
            line(s, x0 + cw / 2, yt, x0 + cw / 2, yb, BLUE)
            line(s, x0 + cw / 2, yb, x0 + cw, yb, BLUE)
            if i < 5:
                line(s, x0 + cw, yb, x0 + cw, yt, BLUE)
    else:
        prev = 0
        for i, b in enumerate(bits):
            x0 = wx + 0.85 + i * cw
            y_ = yt if b else yb
            if i and prev != b:
                line(s, x0, yb if prev == 0 else yt, x0, y_, BLUE)
            line(s, x0, y_, x0 + cw, y_, BLUE)
            prev = b
box(s, wx + 0.85 + 3 * cw + cw / 2 - 0.13, base + 0.9 - 0.5 - 0.13, 0.26, 0.26,
    fill=AMBER, shape=MSO_SHAPE.OVAL)
text(s, 6.8, 5.7, 5.6, 0.7, "Amber dot: cycle 3 is the only cycle where VALID and READY are both high, "
     "so that is where the transfer happens.", size=14, color=MUTED)
footer(s)

# =========================================================================
# 5  Burst
# =========================================================================
s = new_slide(LIGHT, "A burst is one address followed by many beats. AxLEN says how many beats minus "
              "one, so 3 means four beats. WLAST and RLAST mark the final beat so the receiver knows "
              "where it ends. WSTRB says which of the four bytes of a word to write; 1111 is the "
              "whole word. The ID is a small tag put on the request and copied onto its response. "
              "In our design the address simply counts up by 4 for each beat. The example is exactly "
              "what the master does: four words starting at address 0x10.")
title(s, "What a burst is made of")
chips = [("LEN", "beats minus 1"), ("LAST", "marks the final beat"),
         ("WSTRB", "which bytes to write"), ("ID", "tag copied to the response")]
cx = MX
for k, v in chips:
    box(s, cx, 1.55, 2.85, 0.95, fill=BLUE_F)
    text(s, cx + 0.2, 1.65, 2.5, 0.4, k, size=18, bold=True, font=MONO, color=BLUE)
    text(s, cx + 0.2, 2.05, 2.5, 0.4, v, size=15)
    cx += 2.85 + 0.17
text(s, MX, 2.8, 11.9, 0.4, "What the master sends: four words starting at address 0x10, with AxLEN = 3",
     size=17, color=MUTED)
beats = [("beat 0", "0x10", "A5A5_0001"), ("beat 1", "0x14", "5A5A_0002"),
         ("beat 2", "0x18", "DEAD_BEEF"), ("beat 3", "0x1C", "0BAD_F00D")]
bw, gap = 2.85, 0.17
x = MX
for i, (nm, a, d) in enumerate(beats):
    last = (i == 3)
    box(s, x, 3.4, bw, 2.3, fill=AMBER_F if last else WHITE)
    text(s, x + 0.28, 3.6, bw - 0.5, 0.4, nm, size=15, color=MUTED)
    text(s, x + 0.28, 4.05, bw - 0.5, 0.5, a, size=26, bold=True, font=MONO)
    text(s, x + 0.28, 4.7, bw - 0.5, 0.4, "0x" + d, size=16, font=MONO)
    text(s, x + 0.28, 5.15, bw - 0.5, 0.4, "WLAST = 1 on this beat" if last else "WLAST = 0",
         size=14, color=AMBER_D if last else MUTED, bold=last)
    x += bw + gap
text(s, MX, 6.05, 11.9, 0.5, "The address goes up by 4 for each beat, because each word is 4 bytes.",
     size=16, color=MUTED, italic=True)
footer(s)

# =========================================================================
# 6  System architecture
# =========================================================================
s = new_slide(LIGHT, "Walk left to right. The yellow box on the left is the input D flip-flops, loaded by "
              "the testbench. The master takes its data from them and starts the transfers when start "
              "goes high. The five channels, with the real signal names above each wire, go straight "
              "across to the slave: there is no interconnect and nothing in between. On the right the "
              "slave has a write engine, a read engine and a 4 KiB memory. Each read beat is stored "
              "in the output D flip-flops, the yellow box on the far right. B and R run the other way.")
title(s, "System architecture")
picture(s, os.path.join(HERE, "axi4_block_diagram.png"), 0.9, 1.5, w=11.5,
        alt="Block diagram: input D flip-flops, axi4_master, the five channels connecting directly to axi4_slave with write engine, read engine and 4 KiB memory, and output D flip-flops that catch the read data.")
text(s, MX, 6.3, 11.9, 0.3,
     "Yellow boxes are the D flip-flops. The master and the slave are connected directly.",
     size=14, color=MUTED, italic=True)
footer(s)

# =========================================================================
# 7  The master
# =========================================================================
s = new_slide(LIGHT, "The master is one small state machine. It sits in S_IDLE until start goes high. "
              "S_AW sends the write address. S_W sends the four words, one per beat, taken from the "
              "input flip-flops, and raises WLAST on the fourth. S_B waits for the one response. S_AR "
              "sends the read address. S_R takes four beats and hands each one to the output "
              "flip-flops. S_DONE raises done and waits for start to drop, so the sequence can run "
              "again. The master itself holds no data.")
title(s, "The master, step by step")
steps7 = [
    ("S_IDLE", WHITE, "Waits for start."),
    ("S_AW", BLUE_F, "Sends the write address."),
    ("S_W", BLUE_F, "Sends the four words from the input DFFs. WLAST on the fourth."),
    ("S_B", BLUE_F, "Waits for the one write response."),
    ("S_AR", GREEN_F, "Sends the read address."),
    ("S_R", GREEN_F, "Hands each beat to the output DFFs. Ends on RLAST."),
    ("S_DONE", AMBER_F, "done = 1 until start drops."),
]
n7 = 7
gap7 = 0.14
bw7 = (W - 2 * MX - (n7 - 1) * gap7) / n7
x = MX
for i, (nm, fill, d) in enumerate(steps7):
    box(s, x, 1.6, bw7, 2.9, fill=fill, line=RULE if fill == WHITE else None)
    text(s, x + 0.12, 1.78, bw7 - 0.24, 0.5, nm, size=16, bold=True, font=MONO)
    text(s, x + 0.12, 2.4, bw7 - 0.24, 2.0, d, size=14)
    if i < n7 - 1:
        arrow(s, x + bw7 + 0.01, 3.05, x + bw7 + gap7 - 0.01, 3.05, color=MUTED, width=2)
    x += bw7 + gap7
box(s, MX, 4.85, 5.85, 1.8, fill=WHITE)
text(s, MX + 0.3, 5.0, 5.3, 0.4, "Data from the input DFFs", size=18, bold=True, font=HEAD)
text(s, MX + 0.3, 5.5, 5.3, 1.1,
     "The address and four words are not in the master. The testbench loads them into the "
     "input DFFs and the master reads them from there.", size=15)
box(s, MX + 6.05, 4.85, 5.85, 1.8, fill=WHITE)
text(s, MX + 6.35, 5.0, 5.3, 0.4, "Start and done", size=18, bold=True, font=HEAD)
text(s, MX + 6.35, 5.5, 5.3, 1.1,
     "Nothing happens until start goes high. done goes high when the read has finished; "
     "drop start to run it again.", size=15)
footer(s)

# =========================================================================
# 8  The D flip-flop stage
# =========================================================================
s = new_slide(LIGHT, "These are the D flip-flops you asked for, one bank at each end. The input bank is "
              "five 32-bit registers: an address and four words. The testbench writes one register per "
              "clock edge using load, load_sel and load_data. The master reads their outputs directly "
              "and does not change them. The output bank is four registers. Each read beat is stored "
              "in the register for its position, using rd_valid and rd_idx. pass is done AND all four "
              "output registers equal the four input registers. Reset clears both banks.")
title(s, "The D flip-flops at both ends")
box(s, MX, 1.6, 3.82, 3.0, fill=AMBER_F)
text(s, MX + 0.3, 1.8, 3.3, 0.5, "Input DFFs", size=24, bold=True, font=HEAD)
text(s, MX + 0.3, 2.45, 3.3, 0.4, "load, load_sel, load_data", size=14, bold=True, font=MONO)
text(s, MX + 0.3, 3.05, 3.3, 1.5,
     "Five registers: an address and four words. One value is stored per clock edge where load = 1.",
     size=16)
box(s, MX + 4.04, 1.6, 3.82, 3.0, fill=BLUE_F)
text(s, MX + 4.34, 1.8, 3.3, 0.5, "Master", size=24, bold=True, font=HEAD)
text(s, MX + 4.34, 2.45, 3.3, 0.4, "start, done", size=14, bold=True, font=MONO)
text(s, MX + 4.34, 3.05, 3.3, 1.5,
     "Reads the values straight from the input DFFs, writes them, reads them back.", size=16)
box(s, MX + 8.08, 1.6, 3.82, 3.0, fill=AMBER_F)
text(s, MX + 8.38, 1.8, 3.3, 0.5, "Output DFFs", size=24, bold=True, font=HEAD)
text(s, MX + 8.38, 2.45, 3.3, 0.4, "rd_valid, rd_idx, rd_data", size=14, bold=True, font=MONO)
text(s, MX + 8.38, 3.05, 3.3, 1.5,
     "Four registers. Each read beat is stored in its own register as it arrives.", size=16)
box(s, MX, 4.85, 5.85, 1.8, fill=WHITE)
text(s, MX + 0.3, 5.0, 5.3, 0.4, "How pass is made", size=18, bold=True, font=HEAD)
text(s, MX + 0.3, 5.5, 5.3, 1.1,
     "pass = done and the four output registers equal the four input registers.", size=15)
box(s, MX + 6.05, 4.85, 5.85, 1.8, fill=WHITE)
text(s, MX + 6.35, 5.0, 5.3, 0.4, "Direct connection", size=18, bold=True, font=HEAD)
text(s, MX + 6.35, 5.5, 5.3, 1.1,
     "Nothing sits between the master and the slave, so a beat moves every clock "
     "the receiver is ready.", size=15)
footer(s)

# =========================================================================
# 9  The slave
# =========================================================================
s = new_slide(LIGHT, "The slave has two state machines because reads and writes use different "
              "channels. The write engine takes the address, then raises WREADY and writes each "
              "word into memory, using WSTRB to pick bytes. When WLAST arrives it sends one response. "
              "The read engine takes the address, then sends the first word, and fetches each next "
              "word as the previous one is taken. Every response is OKAY: this slave does no request "
              "checking, so an address beyond the 4 KiB simply wraps inside the memory.")
title(s, "The slave")
card_row(s, 1.6, 4.9, [
    (VIOLET_F, "Write engine", "IDLE, DATA, RESP",
     "Takes the address. Then takes beats until WLAST, writing each word and adding 4 to the address. "
     "WSTRB picks the bytes. Then sends one response."),
    (GREEN_F, "Read engine", "IDLE, DATA",
     "Takes the address. Sends ARLEN + 1 beats, adding 4 each time, and marks the last with RLAST."),
    (WHITE, "Responses", "",
     "BRESP and RRESP are always 00, OKAY. There is no request checking. "
     "An address beyond the 4 KiB wraps inside the memory."),
], 3, body_size=17)
footer(s)

# =========================================================================
# 10  One sequence end to end
# =========================================================================
s = new_slide(LIGHT, "The whole sequence, with the real cycle numbers from the simulation. Cycle 0 is the "
              "first cycle with start high. The write address is taken at cycle 1: the slave was "
              "already ready, so there is no wait. The four data beats are taken at cycles 2, 3, 4 and "
              "5, one every clock, and WLAST is on the fourth. The response is taken at cycle 6. The "
              "read address is taken at cycle 7, the four beats arrive at cycles 8 to 11, and done "
              "goes high at cycle 12. That is 13 clock cycles from start.")
title(s, "One sequence, end to end")
text(s, MX, 1.4, 11.9, 0.5, "The master writes four words, then reads them back.",
     size=20, color=MUTED)
steps10 = [
    ("1  Write address", BLUE_F, "AW taken at cycle 1. The slave is ready, so no wait."),
    ("2  Write data", BLUE_F, "Four beats taken at cycles 2, 3, 4, 5. WLAST on the fourth."),
    ("3  Response", VIOLET_F, "B taken at cycle 6, right after the last beat."),
    ("4  Read", GREEN_F, "AR taken at cycle 7, then four beats at cycles 8 to 11. RLAST on the fourth."),
    ("5  Done", AMBER_F, "done = 1 at cycle 12, pass = 1. 13 cycles from start."),
]
bw, gap = 2.18, 0.25
x = MX
for i, (hd, fill, d) in enumerate(steps10):
    box(s, x, 2.25, bw, 3.7, fill=fill)
    text(s, x + 0.2, 2.45, bw - 0.4, 0.9, hd, size=20, bold=True, font=HEAD)
    text(s, x + 0.2, 3.45, bw - 0.4, 2.4, d, size=16)
    if i < 4:
        arrow(s, x + bw + 0.02, 4.1, x + bw + gap - 0.02, 4.1, color=MUTED, width=2)
    x += bw + gap
text(s, MX, 6.3, 11.9, 0.5, "The next two slides show the write and the read as waveforms.",
     size=16, color=MUTED, italic=True)
footer(s)

# =========================================================================
# 11  Write waveform
# =========================================================================
s = new_slide(LIGHT, "This waveform was recorded from our simulation. Amber dots are handshakes. The "
              "master raises AWVALID at cycle 1 and AWREADY is already high, so the address is taken "
              "at once. The four data beats are taken at cycles 2, 3, 4 and 5: WREADY stays high, so "
              "one beat moves every clock. WLAST is high on the fourth. BVALID comes at cycle 6 and "
              "BREADY is already high, so the response is taken immediately.")
title(s, "Write: four words, recorded from the simulation")
pw = 4.3 * 1944 / 846
picture(s, os.path.join(HERE, "axi4_waveform_write.png"), (W - pw) / 2, 1.45, h=4.3,
        alt="Waveform of the master writing four words straight to the slave.")
cal = [
    ("Amber dots", "are handshakes."),
    ("Address", "taken at 1: the slave was already ready."),
    ("Data beats", "taken at 2, 3, 4, 5: one every clock."),
    ("Response", "taken at 6, right after WLAST."),
]
cw = (W - 2 * MX - 3 * 0.2) / 4
x = MX
for hd, bd in cal:
    box(s, x, 5.85, cw, 1.0, fill=WHITE)
    text(s, x + 0.2, 5.95, cw - 0.4, 0.9, [{"text": hd, "bold": True, "size": 15, "after": 2},
                                           {"text": bd, "size": 14, "color": MUTED}])
    x += cw + 0.2
footer(s)

# =========================================================================
# 12  Read waveform
# =========================================================================
s = new_slide(LIGHT, "Same four words read back. The read address is taken at cycle 1 of this window, "
              "which is cycle 7 of the whole run. The four beats come back at cycles 2, 3, 4 and 5 of "
              "the window, one every clock, with RLAST on the fourth. Then done goes high at cycle 6. "
              "Each beat is stored in the output D flip-flops as it arrives, and pass says they equal "
              "what was written.")
title(s, "Read: the same four words, from the simulation")
pw = 3.8 * 1944 / 750
picture(s, os.path.join(HERE, "axi4_waveform_read.png"), (W - pw) / 2, 1.6, h=3.8,
        alt="Waveform of the master reading four words back from the slave.")
cal = [
    ("Address", "taken at 1: the slave was already ready."),
    ("Data beats", "taken at 2, 3, 4, 5: one every clock."),
    ("RLAST", "rides on the fourth beat, at cycle 5."),
    ("done", "goes high at 6. The sequence took 13 cycles."),
]
cw = (W - 2 * MX - 3 * 0.2) / 4
x = MX
for hd, bd in cal:
    box(s, x, 5.55, cw, 1.2, fill=WHITE)
    text(s, x + 0.2, 5.67, cw - 0.4, 1.0, [{"text": hd, "bold": True, "size": 15, "after": 2},
                                           {"text": bd, "size": 14, "color": MUTED}])
    x += cw + 0.2
footer(s)

# =========================================================================
# 13  Verification
# =========================================================================
s = new_slide(LIGHT, "We test two things. First the slave on its own: the testbench plays the master "
              "and drives each channel by hand, so it can force awkward cases such as data before "
              "address, held-off responses, byte strobes, the longest burst of 256 beats, and a "
              "reset in the middle of a burst. Second the whole system: the testbench loads the "
              "input flip-flops, raises start, and checks that the master finishes, that pass is set, "
              "that the output flip-flops and the slave's memory hold the four words, that every "
              "response is OKAY with the right ID, and that VALID and the payload hold steady until "
              "READY. It also runs a second set of values at another address, and resets "
              "mid-sequence.")
title(s, "How it was verified")
card_row(s, 1.55, 4.4, [
    (BLUE_F, "Slave on its own", "tb_axi4_slave.v",
     "The testbench plays the master. 33 checks: single and multi-beat bursts, data first or address first, "
     "held-off responses, byte strobes, ID echo, 256 beats, reset mid-burst."),
    (GREEN_F, "Whole system", "tb_axi4_top.v",
     "Input DFFs, master, slave, output DFFs. 43 checks: idle until start, words in memory and in the "
     "output DFFs, pass, OKAY responses, VALID holds, a second run with other values, reset mid-sequence."),
], 2, body_size=15)
box(s, MX, 6.15, 11.9, 0.62, fill=AMBER_F)
text(s, MX + 0.3, 6.15, 11.3, 0.62, "Total: 33 + 43 = 76 checks, 0 failing.   Run it:  .\\run_tests.ps1",
     size=17, bold=True, anchor=MSO_ANCHOR.MIDDLE)
footer(s)

# =========================================================================
# 14  Limits
# =========================================================================
s = new_slide(LIGHT, "Say these before you are asked. The burst shape is fixed: the master always writes "
              "and reads four words with all bytes enabled. The input flip-flops must not change "
              "while a sequence runs, because the master reads them live. Only counting-up bursts, so "
              "no FIXED or WRAP and no narrow beats. Every response is OKAY and an out-of-range address "
              "wraps. One burst at a time. And the lock, cache, protection and QoS signals are not "
              "included.")
title(s, "Limits, stated plainly")
lims = [
    ("Fixed burst shape", "The master always writes and reads four words, with all four bytes enabled."),
    ("Load, then start", "The input DFFs must not change while a sequence runs. The master reads them live."),
    ("Counting-up bursts only", "No FIXED or WRAP bursts and no narrow beats. The address adds 4 per beat."),
    ("Always OKAY", "No error responses. An address beyond the 4 KiB wraps inside the memory."),
    ("One burst at a time", "One write or one read in progress. IDs are echoed, never used to reorder."),
    ("No attribute signals", "AxLOCK, AxCACHE, AxPROT and AxQOS are not included."),
]
for i, (h_, d) in enumerate(lims):
    col, row = i % 2, i // 2
    x = MX + col * 6.05
    y = 1.55 + row * 1.7
    box(s, x, y, 5.85, 1.5, fill=WHITE)
    text(s, x + 0.3, y + 0.18, 5.3, 0.4, h_, size=18, bold=True, font=HEAD)
    text(s, x + 0.3, y + 0.68, 5.3, 0.75, d, size=14, color=MUTED)
footer(s)

# =========================================================================
# 15  Demo
# =========================================================================
s = new_slide(NAVY, "Live demo, about two minutes. Open PowerShell in E:\\FPGA_Project and run "
              ".\\run_tests.ps1. It compiles and runs two suites. Point at the two TEST PASSED lines "
              "and the check counts, 33 and 43, and the line that says the whole sequence took 13 "
              "clock cycles. Then run the top suite with -Wave to open the waveform in GTKWave and "
              "show the same shape as the recorded slides: the handshakes, with WREADY high so a "
              "beat moves every clock. If the live run fails for any reason, slides 11 and 12 are "
              "recorded from this exact simulation.")
title(s, "Live demo", dark=True)
box(s, MX, 1.6, 7.4, 3.0, fill="0F1726")
text(s, MX + 0.4, 1.85, 6.8, 2.6, [
    {"text": "cd E:\\FPGA_Project", "font": MONO, "size": 18, "color": ON_DARK, "after": 14},
    {"text": ".\\run_tests.ps1", "font": MONO, "size": 18, "color": AMBER, "after": 14},
    {"text": ".\\run_tests.ps1 -Suite top -Wave", "font": MONO, "size": 18, "color": AMBER, "after": 0}])
text(s, MX + 8.0, 1.6, 3.9, 0.4, "What to point at", size=20, bold=True, color=AMBER, font=HEAD)
text(s, MX + 8.0, 2.15, 3.9, 3.0, [
    {"text": "Two suites, each ending TEST PASSED", "after": 10},
    {"text": "Check counts: 33 and 43", "after": 10},
    {"text": "\"the whole sequence took 13 clock cycles\"", "after": 0}],
     size=16, color=ON_DARK)
box(s, MX, 5.0, 11.9, 1.5, fill="22304A")
text(s, MX + 0.4, 5.2, 11.1, 1.2, [
    {"text": "If the live run fails", "size": 17, "bold": True, "color": AMBER, "after": 6},
    {"text": "Slides 11 and 12 are recorded from this same simulation. Logs and waveforms are "
             "written to the build folder.", "size": 16, "color": ON_DARK}])
footer(s, dark=True)

# =========================================================================
# 16  Questions
# =========================================================================
s = new_slide(LIGHT, "Likely questions with short answers. Why five channels: address, data and "
              "response move independently. Is there an interconnect: no, one master and one slave "
              "are wired directly. Where does the data come from: the input D flip-flops, loaded by "
              "the testbench. Where does the read data go: the output D flip-flops, one register per "
              "beat. How do we know it worked: the master's pass compares them, and the testbench "
              "checks the slave's memory. How does the slave know the burst ended: WLAST on writes, "
              "a beat count on reads.")
title(s, "Questions we expect")
qa = [
    ("Why five channels?", "Address, data and response move independently, so a slow response never blocks the next address."),
    ("Is there an interconnect?", "No. One master and one slave are wired directly. An interconnect is for more of either."),
    ("Where does the data come from?", "From the input D flip-flops, loaded by the testbench. The master reads them and writes to the slave."),
    ("Where does the read data go?", "Into the output D flip-flops, one register per beat. pass compares them with the input."),
    ("How do we know it worked?", "pass compares read-back with what was written. The testbench also checks the slave's memory."),
    ("How does the slave know the end?", "On writes, from WLAST. On reads, by counting ARLEN + 1 beats."),
]
for i, (q, a) in enumerate(qa):
    col, row = i % 2, i // 2
    x = MX + col * 6.05
    y = 1.55 + row * 1.72
    box(s, x, y, 5.85, 1.55, fill=WHITE)
    text(s, x + 0.3, y + 0.17, 5.3, 0.4, q, size=17, bold=True, font=HEAD)
    text(s, x + 0.3, y + 0.66, 5.3, 0.85, a, size=15, color=MUTED)
footer(s)


prs.save(OUT)
print("saved", OUT, "slides:", len(prs.slides))
